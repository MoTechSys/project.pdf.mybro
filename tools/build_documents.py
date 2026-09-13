#!/usr/bin/env python3
"""Generate editable diagrams, workbook and an Arabic report from campus.json."""
from pathlib import Path
import csv, html, json, math, xml.etree.ElementTree as ET
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
ROOT=Path(__file__).resolve().parents[1]
m=json.loads((ROOT/'design/campus.json').read_text()); ss=m['subnets']; services=m['services']
validation=json.loads((ROOT/'tests/validation_results.json').read_text())
runtime_path=ROOT/'tests/runtime_results.json'
runtime=json.loads(runtime_path.read_text()) if runtime_path.exists() else {}
for directory in ('diagrams','report','pt','screenshots','tests'): (ROOT/directory).mkdir(exist_ok=True)

def out(path,text): (ROOT/path).write_text(text,encoding='utf-8')
def mdtable(headers,rows): return '| '+' | '.join(map(str,headers))+' |\n|'+'|'.join(['---']*len(headers))+'|\n'+'\n'.join('| '+' | '.join(str(v).replace('\n',' / ') for v in row)+' |' for row in rows)+'\n'

# Workbook: all design tables and explicit evidence status, no hidden assumptions.
wb=Workbook(); wb.remove(wb.active)
for filename,title in [('addressing_plan','VLSM'),('vlan_matrix','VLAN Matrix'),('port_map','Port Map'),('devices','Devices'),('endpoints','Endpoints')]:
 ws=wb.create_sheet(title)
 with (ROOT/f'design/{filename}.csv').open(encoding='utf-8-sig') as f:
  for row in csv.reader(f): ws.append(row)
 ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]: cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='17324D')
 for col in ws.columns:
  ws.column_dimensions[col[0].column_letter].width=min(36,max(14,max(len(str(c.value or '')) for c in col)+2))
 for row in ws.iter_rows(min_row=2):
  for c in row: c.alignment=Alignment(vertical='top'); c.fill=PatternFill('solid',fgColor='F0F5F8' if c.row%2==0 else 'FFFFFF')
ws=wb.create_sheet('Assumptions');
for row in [('Item','Value'),('Counts','Assumed concurrent endpoints, not supplied by PDF'),('Growth','25%'),('Reserved per subnet','15 usable addresses; first 3 for HSRP'),('Simulation','46 representative endpoints, not every planned user'),('Runtime evidence','See tests/runtime_results.json; partial verification, not ready for final submission'),('Static checks',str(validation['passed'])+' passed; '+str(validation['failed'])+' failed'),('DHCP redundancy','Single server; HSRP does not replicate leases'),('Internet','Documentation address space and simulated ISP'),('Lab credentials','Disposable examples only; replace before real deployment')]: ws.append(row)
ws.column_dimensions['A'].width=25; ws.column_dimensions['B'].width=95
wb.save(ROOT/'design/addressing_plan.xlsx')

# Clean overview diagram. The full port map is authoritative; overview aggregates access endpoints.
W,H=2600,1580
im=Image.new('RGB',(W,H),'#f2f6fa'); d=ImageDraw.Draw(im)
font='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'; bold='/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
F=lambda n,b=False: ImageFont.truetype(bold if b else font,n)
d.rectangle((0,0,W,150),fill='#142d46'); d.text((65,32),'UNIVERSITY CAMPUS NETWORK',font=F(46,True),fill='white'); d.text((68,97),'5 buildings | hierarchical architecture | OSPF Area 0 | HSRP | segmented access',font=F(24),fill='#b9d8e9')
positions={'PUBLIC-WEB':(1300,200),'ISP':(1300,295),'EDGE':(1300,390),'CORE-1':(830,520),'CORE-2':(1770,520)}
for bi,b in enumerate('ABCDE'):
 x=300+bi*500
 positions.update({f'DIST-{b}1':(x-112,815),f'DIST-{b}2':(x+112,815),f'ACC-{b}1':(x-112,1045),f'ACC-{b}2':(x+112,1045)})
 d.rounded_rectangle((x-228,688,x+228,1337),radius=20,fill='white',outline='#c9d7e2',width=2)
 d.text((x-205,709),f'BUILDING {b}',font=F(23,True),fill='#17324d')
 d.text((x-205,743),m['buildings'][b]['name'],font=F(17),fill='#596c7d')
 d.text((x-205,1150),m['buildings'][b]['allocation'],font=F(24,True),fill='#116b72')
 vs=[str(s['vlan']) for s in ss if s['building']==b]
 d.text((x-205,1194),'VLANs: '+', '.join(vs[:6]),font=F(17),fill='#445566'); d.text((x-205,1223),', '.join(vs[6:]),font=F(17),fill='#445566')
 d.text((x-205,1270),'Local L2 domains / routed uplinks',font=F(16),fill='#667788')
for l in m['links']:
 if l['a'] not in positions or l['b'] not in positions: continue
 a=positions[l['a']]; b=positions[l['b']]
 color='#187895' if l['kind']=='routed' else '#d99232'
 d.line((a,b),fill=color,width=4)
for name,(x,y) in positions.items():
 small=name.startswith(('DIST','ACC')); w=195 if small else 250; h=74
 color='#157f84' if name.startswith('DIST') else '#566b80' if name.startswith('ACC') else '#17324d'
 d.rounded_rectangle((x-w//2,y-h//2,x+w//2,y+h//2),radius=12,fill=color)
 tw=d.textbbox((0,0),name,font=F(22,True))[2]; d.text((x-tw/2,y-15),name,font=F(22,True),fill='white')
d.text((65,1370),'Core: 2 x 3650   |   Distribution: 10 x 2911   |   Access: 10 x 2960',font=F(27,True),fill='#17324d')
d.text((65,1423),'E services: DHCP-1 / DNS / educational web / HR / Finance / IT management',font=F(24),fill='#445566')
d.text((65,1470),'Blue = routed link   Gold = trunk   DIST-X1 = HSRP active   DIST-X2 = standby   ACC-X1 = STP root',font=F(23),fill='#445566')
d.text((65,1520),'DESIGN DIAGRAM — not a Packet Tracer screenshot. See design/port_map.csv for every cable.',font=F(21),fill='#985814')
im.save(ROOT/'diagrams/topology.png')
# Editable diagrams.net topology, one vertex per infrastructure or end device and one edge per cable.
mx=ET.Element('mxfile',host='app.diagrams.net'); diagram=ET.SubElement(mx,'diagram',name='Campus topology',id='campus')
model=ET.SubElement(diagram,'mxGraphModel',dx='2600',dy='1800',grid='1',page='1',pageWidth='2800',pageHeight='2300'); root=ET.SubElement(model,'root')
ET.SubElement(root,'mxCell',id='0'); ET.SubElement(root,'mxCell',id='1',parent='0')
coords={k:(x-90,y-35) for k,(x,y) in positions.items()}
coords['DHCP-1']=(2330,1440)
for b in 'ABCDE':
 ends=[e for e in m['endpoints'] if e['building']==b]
 for j,e in enumerate(ends): coords[e['name']]=(90+'ABCDE'.index(b)*500+(j%2)*225,1440+(j//2)*95)
for node in m['nodes']+m['endpoints']:
 name=node['name']; x,y=coords.get(name,(50,2100)); color='#17324D' if name.startswith('CORE') else '#147E83' if name.startswith('DIST') else '#566B80'
 cell=ET.SubElement(root,'mxCell',id=name,value=name,style=f'rounded=1;whiteSpace=wrap;html=1;fillColor={color};fontColor=#ffffff;strokeColor=none;',vertex='1',parent='1')
 ET.SubElement(cell,'mxGeometry',x=str(x),y=str(y),width='180',height='65',attrib={'as':'geometry'})
for i,l in enumerate(m['links']):
 cell=ET.SubElement(root,'mxCell',id=f'link{i}',value=f'{l["a_port"]} / {l["b_port"]}',style='edgeStyle=orthogonalEdgeStyle;endArrow=none;html=1;fontSize=8;strokeColor='+('#187895' if l['kind']=='routed' else '#d99232')+';',edge='1',parent='1',source=l['a'],target=l['b'])
 ET.SubElement(cell,'mxGeometry',relative='1',attrib={'as':'geometry'})
ET.ElementTree(mx).write(ROOT/'diagrams/topology.drawio',encoding='utf-8',xml_declaration=True)

# Report content is shared by Markdown and DOCX so technical claims stay synchronized.
sections=[]
def heading(t): sections.append(('h',t))
def paragraph(t): sections.append(('p',t))
def table(headers,rows): sections.append(('t',(headers,rows)))
def subheading(t): sections.append(('h2',t))
def code(t): sections.append(('code',t))
def config(name): return (ROOT/f'configs/{name}.txt').read_text()
def block(name,opening):
 lines=config(name).splitlines(); start=lines.index(opening)
 end=next(i for i in range(start+1,len(lines)) if lines[i]=='exit')
 return '\n'.join(lines[start:end+1])

heading('ملخص تنفيذي وحالة التسليم')
paragraph('تصميم شبكة جامعة من خمسة مبانٍ وأربع كليات، مع مركز بيانات مستقل وإدارة شبكية محمية. تم إعداد التصميم والعنونة والإعدادات من نموذج بيانات موحد. أعداد الأجهزة افتراضات تصميمية وليست أرقامًا واردة في التكليف. لا تُعد الفحوص البرمجية بديلًا عن تشغيل المحاكي.')
paragraph(f'حالة الأدلة عند توليد هذا التقرير: نجح {validation["passed"]} فحصًا ساكنًا وفشل {validation["failed"]}. ملف university.pkt مبني ومحفوظ بالمحاكي؛ سجل الاختبارات التشغيلية يتضمن {sum(v.get("status")=="PASS" for v in runtime.values())} اختبارات ناجحة موثقة. الاختبارات غير المنفذة موضحة بصراحة ولا تعد ناجحة ضمنيًا. المرجع التفصيلي tests/runtime_results.json.')
heading('الفصل الأول: تحليل الشبكة')
paragraph('المصدر المرجعي هو project.pdf، التكليف النهائي العملي لمقرر تحليل وتصميم الشبكات، 14 صفحة. يُعتمد عنوان كل متطلب بدل الترقيم غير المتسلسل في المستند. لم يحدد الملف برنامج المحاكاة ولا أعداد المستخدمين ولا ميزانية الأجهزة.')
table(['المبنى','الجهة','كتلة العناوين','نقاط النهاية المفترضة'],[[b,v['name'],v['allocation'],sum(s['assumed_endpoints'] for s in ss if s['building']==b)] for b,v in m['buildings'].items()])
paragraph('الأقسام: A الأمن السيبراني والشبكات والبرمجيات والمختبرات والإدارة؛ B المدنية والكهربائية والمعامل والإدارة؛ C الفيزياء والكيمياء والمختبرات والإدارة؛ D إدارة الأعمال والمحاسبة والاقتصاد والإدارة؛ E إدارة الجامعة وشؤون الطلاب والموارد البشرية والمالية ومركز البيانات وإدارة الشبكة.')
paragraph('الأعداد تمثل نقاط اتصال متزامنة: حواسيب مستخدمين وأجهزة مختبرات وضيوف وهواتف وكاميرات وخوادم وأجهزة إدارة؛ لا تمثل عدد أشخاص فريدًا لأن الشخص قد يستخدم أكثر من جهاز. يضاف نمو 25% و15 عنوانًا محجوزًا لكل شبكة. الهواتف هنا شبكة بيانات مخصصة؛ إعداد المكالمات وCME ليس من متطلبات الملف ولا يُدّعى تنفيذه.')
paragraph('المتطلبات الوظيفية: DHCP، DNS وخدمة ويب تعليمية، OSPF، HSRP، ACL، SSH، VLAN 99 للإدارة، شبكة خوادم مستقلة، Port Security وDHCP Snooping. المتطلبات غير الوظيفية: تقليل نطاقات البث، عزل الأعطال، سهولة التوسع، وضوح التوثيق وقابلية إعادة توليد الإعدادات.')
heading('الفصل الثاني: تصميم الشبكة')
paragraph('طبقة Core تتكون من زوج 3650. طبقة Distribution تضم زوج راوترات 2911 لكل مبنى، مسؤولًا عن التوجيه عبر واجهات فرعية 802.1Q وHSRP وACLs. لكل راوتر وصلة موجهة /30 إلى كل Core ووصلة Trunk إلى سويتش الوصول المحلي. يرتبط زوج سويتشات 2960 داخل المبنى بوصلة Trunk لتبادل VLANs والوصول للبوابة الاحتياطية. لا تمتد Layer 2 بين المباني.')
paragraph('سبب اختيار راوترات التوزيع: أظهر الفحص داخل Packet Tracer 9 فقد ارتباط ACL بواجهات SVI على سويتشات Layer 3 بعد الحفظ، بينما احتفظت واجهات 2911 الفرعية بقواعد ip access-group. تم اعتماد الحل الذي يحافظ على العزل بعد إعادة الفتح بدل الاعتماد على إعدادات نصية لا ينفذها المحاكي. هذا اختيار للمختبر؛ قد يكون زوج سويتشات Layer 3 الحقيقية أعلى أداء في الإنتاج.')
paragraph('التصميم الفيزيائي المقترح: Core في غرفة مركز البيانات بالمبنى E، والتوزيع في غرفة اتصالات كل مبنى، والوصول في خزائن الطوابق. الربط البعيد المقترح بالألياف وفق المسافة وميزانية البصريات، وربط الأجهزة النهائية بالنحاس ضمن حدود المسافة المناسبة. المحاكاة التعليمية تمثل الوظائف ولا تثبت سعة الألياف أو الأداء الواقعي.')
paragraph('عدد سويتشات Access في ملف المحاكاة لا يكفي لتوصيل جميع أعداد التخطيط الفعلية؛ الأجهزة الطرفية عينات تمثيلية. لتقدير الحجم الفعلي يُحسب عدد منافذ الوصول من مجموع نقاط النهاية بعد النمو، مع مراعاة المنافذ المحجوزة ومتطلبات PoE. الجدول أدناه حد أدنى تقريبي لسويتشات 24 منفذًا بافتراض منفذ مستقل لكل نقطة نهاية، وليس قائمة شراء نهائية.')
table(['المبنى','نقاط النهاية بعد النمو','حد أدنى لسويتشات 24 منفذًا','الممثل في المحاكاة'],[[b,sum(s['growth_endpoints'] for s in ss if s['building']==b),math.ceil(sum(s['growth_endpoints'] for s in ss if s['building']==b)/24),2] for b in m['buildings']])
paragraph('قرار VLAN IDs: تُستخدم أرقام متسقة للأدوار العامة مثل 10 و20 و60 و99، مع أرقام إضافية لعزل الأقسام 31 و32 والجهات الحساسة 11 و12 و13. VLAN 30 تشير إلى أول شبكة طلاب بالقسم المحدد في اسمها؛ أسماء VLANs وعناوين الشبكات تزيل الالتباس. تكرار الرقم لا يدمج نطاقات البث ما دام الربط بين المباني Layer 3.')
table(['معيار المقارنة','تبرير القرار'],[['الأمان','العزل ينتج من حدود Layer 3 والسياسة، وليس اختلاف الرقم وحده.'],['التوسع','يمكن إضافة قسم وشبكة داخل كتلة المبنى دون مد Layer 2 إلى Core.'],['الإدارة','توحيد أرقام الأدوار المتكررة يقلل الأخطاء؛ اسم المبنى والقسم إلزامي في الجداول.'],['Broadcast Domains','لكل مبنى وشبكة نطاق بث مستقل، حتى عند تكرار VLAN ID.'],['استكشاف الأعطال','تحديد المبنى ثم VLAN ثم Subnet يقصر نطاق التشخيص.']])
paragraph('VLAN 50 موجودة في E فقط للخدمات المركزية. VLAN 99 محلية لكل مبنى، مع جهاز IT معتمد في E. VLAN 998 Native غير مستخدمة للأجهزة، وVLAN 999 للمنافذ المغلقة ولا تملك واجهة Layer 3 الفرعية. شبكات الأقسام موضحة كاملة في جدول VLAN Matrix في ملف Excel.')
paragraph('DIST-X1 أولوية HSRP 110 وDIST-X2 أولوية 100. على سويتشات الوصول: ACC-X1 أولوية STP 4096 وACC-X2 أولوية 8192 مع Rapid-PVST. الراوتر لا يشارك في STP. لا يوجد PortFast أو BPDU Guard على رابطَي Trunk. تعطل راوتر البوابة يمكن تجاوزه عبر وصلة الوصول البينية، لكن تعطل سويتش الوصول يفصل الأجهزة الموصولة به؛ هذا حد معلن للنموذج.')
heading('الفصل الثالث: خطة العنونة وVLSM')
paragraph('لكل شبكة، المطلوب = سقف(عدد الأجهزة × 1.25) + 15 عنوانًا محجوزًا. يُختار أصغر حجم كتلة يحقق 2^h - 2 ≥ المطلوب، ثم القناع /(32-h). تُرتب الشبكات من الأكبر إلى الأصغر داخل /16 الخاص بالمبنى. هذا يحقق VLSM دون تداخل.')
paragraph('العنوان الأول هو بوابة HSRP الافتراضية، والثاني لـDIST-X1، والثالث لـDIST-X2. العناوين من +1 إلى +15 محجوزة. يبدأ DHCP من +16 إلى آخر عنوان مضيف. شبكات الإدارة والخوادم والكاميرات بعناوين ثابتة؛ الباقي يستخدم DHCP.')
table(['المبنى/VLAN','الأجهزة','بعد النمو','الشبكة','السعة القابلة للاستخدام','سعة الأجهزة بعد الحجز'],[[f'{s["building"]}/{s["vlan"]}',s['assumed_endpoints'],s['growth_endpoints'],s['network'],s['usable'],s['endpoint_capacity']] for s in ss])
paragraph('جدول الحدود التالي يكمل عناصر التكليف: أول وآخر مضيف، Broadcast، والقناع. البوابة الافتراضية تساوي أول مضيف في كل صف. جميع حقول التفاصيل، بما فيها عناوين جهازي HSRP ومدى DHCP، موجودة أيضًا في addressing_plan.xlsx وCSV.')
table(['المبنى/VLAN','القناع','أول مضيف / Gateway','آخر مضيف','Broadcast'],[[f'{s["building"]}/{s["vlan"]}',s['mask'],s['first_host'],s['last_host'],s['broadcast']] for s in ss])
paragraph('روابط النقل تُقسم من 10.0.0.0/24 إلى /30، وRouter IDs عناوين Loopback /32 من 10.255.255.0/24. تبقى المساحات غير المخصصة داخل /16 احتياطًا للتوسع. لا نعلن تجميع /16 وهميًا؛ OSPF Area 0 يعلن الشبكات الفعلية، ويمكن تصميم مناطق وتلخيص حدودي لاحقًا عند توسع مبرر.')
heading('الفصل الرابع: إعداد الشبكة وتنفيذها')
paragraph('ترتيب التطبيق: وضع الأجهزة بالموديلات المحددة، توصيل port_map.csv، إعداد Core ثم Distribution ثم Access، ثم DHCP-1 وEDGE وISP، ثم ضبط الخوادم وأجهزة المستخدمين. تُلصق ملفات configs من وضع User EXEC؛ إذا ظهرت رسالة Initial configuration dialog يكون الجواب no. يجب مراجعة رسائل رفض الأوامر وعدم تجاهلها.')
paragraph('الموديلات المستهدفة 3650-24PS و2960-24TT و2911. أسماء المنافذ في الجداول جزء من التصميم. دعم صياغة بعض الأوامر مثل RSA وHSRP track يتطلب التحقق على إصدار Packet Tracer الفعلي؛ استخدام جهاز مختلف دون تعديل المنافذ ليس إجراء صحيحًا.')
paragraph('Inter-VLAN Routing يحدث على راوترات التوزيع عبر GigabitEthernet0/2.<VLAN> باستخدام encapsulation dot1Q. الجهاز يرسل إلى MAC بوابة HSRP؛ الراوتر Active يفحص ACL الواردة على الواجهة الفرعية ويبحث عن الوجهة المتصلة أو مسار OSPF. حركة VLANs داخل المبنى توجه محليًا دون إلزامها بالمرور عبر Core. VLAN 998 Native غير مستخدمة لحركة المستخدمين.')
paragraph(f'DHCP مركزي على راوتر خدمة DHCP-1 بعنوان {services["dhcp"]} داخل VLAN 50. يوفر أوامر show ip dhcp binding المطلوبة في التكليف. كل واجهة Layer 3 الفرعية ديناميكية في كلا جهازي التوزيع تحمل ip helper-address. DHCP يوزع VIP وليس العنوان الحقيقي لأحد جهازي التوزيع. الخيار مركزي وبسيط لكنه نقطة فشل منفردة للخدمة؛ HSRP لا ينسخ leases ولا يحقق DHCP redundancy.')
table(['الخدمة','العنوان','الإعداد'],[['DHCP-1',services['dhcp'],'Scopes في configs/DHCP-1.txt'],['DNS',services['dns'],'تفعيل DNS وإنشاء سجلات edu.campus.example وwww.external.example'],['Educational Web',services['educational_web'],'تفعيل HTTP وHTTPS إن كان مدعومًا'],['IT-ADMIN',services['it_admin'],'Static / VLAN 99 / المصدر الوحيد لإدارة SSH'],['PUBLIC-WEB',services['public_web'],'خادم خارجي محاكى؛ ليس إنترنت حقيقيًا']])
paragraph('OSPF process 1 وArea 0 واحدة: اختيار مناسب لحجم النموذج ويقلل التعقيد. Router ID فريد لكل Core وDistribution وEDGE. passive-interface default ثم تفعيل الجيران على الروابط الموجهة فقط. واجهات Layer 3 الفرعية تُعلن كشبكات دون إنشاء جيران مع المستخدمين. الروابط الموجهة point-to-point، لذلك لا نعتمد انتخاب DR/BDR عليها.')
paragraph('عدد الجيران المتوقع في الحالة المستقرة: لكل DIST جارَان من Core؛ لكل Core عشرة أجهزة Distribution وCore الآخر وEDGE، أي 12 جارًا؛ EDGE له جاران. هذه أعداد متوقعة في التصميم وليست مخرجات أوامر فعلية.')
paragraph('HSRP group يساوي VLAN ID وVIP أول مضيف؛ X1 Active افتراضيًا وX2 Standby، مع preempt لإعادة الدور بعد استعادة الجهاز. الاختبار المعتمد هو إيقاف الواجهة LAN للراوتر Active أو الجهاز نفسه. لم تحتفظ المحاكاة بأوامر تتبع الوصلات في التجربة الأولية؛ لذلك أُزيلت من النسخة المعتمدة ولا يُدّعى أن فقد رابطَي Core يغيّر أولوية HSRP. فقد رابط Core واحد يعالج بمسار OSPF الآخر؛ فقدهما معًا دون تعطل LAN قيد يحتاج Object Tracking على منصة تدعمه.')
paragraph('الإنترنت ممثل بخادم 203.0.113.10 خلف ISP، وربط EDGE إلى ISP بالشبكة 198.51.100.0/30. هذه عناوين محجوزة للتوثيق. EDGE ينفذ PAT وله مسار افتراضي يعلنه في OSPF. EDGE وISP والخادم الخارجي نقاط فشل منفردة معلنة؛ التكليف يطلب Gateway Redundancy ولا يساوي ذلك ضمان عدم وجود أي نقطة فشل في الجامعة.')
subheading('دليل التنفيذ التفصيلي: كيف تستخدم الأوامر التالية؟')
paragraph('إذا فتحت university.pkt فلا تعِد لصق الإعدادات تلقائيًا؛ افحص الموجود أولًا. الخطوات التالية تشرح البناء من الصفر أو مراجعة الإعداد. الأوامر مقتطفة من ملفات configs الفعلية، وليست مخرجات اختبارات. لكل جهاز ملف مستقل كامل؛ كما أُرفق configuration_reference.docx بجميع إعدادات الأجهزة الخمسة والعشرين دون حذف. لا تلصق ملف جهاز على جهاز آخر، ولا تنسخ جميع الأجهزة إلى جلسة واحدة.')
paragraph('لغة الشرح والجداول عربية باتجاه اليمين إلى اليسار. أوامر IOS وأسماء المنافذ والعناوين تظل باتجاه لاتيني صحيح حتى لا تنعكس عند النسخ. السطر enable ينقل إلى Privileged EXEC، وconfigure terminal إلى وضع الإعداد العام، وinterface إلى إعداد منفذ، وexit يرجع مستوى واحدًا، وend يعود إلى EXEC. لا تكتب الرموز Router# أو Router(config)# ضمن الأمر.')
subheading('الخطوة ١ — وضع الأجهزة وتوصيل الكابلات')
paragraph('من Network Devices اختر Switches ثم ضع Core وAccess بالموديلات المذكورة؛ ومن Routers ضع 2911 للتوزيع وEDGE وISP وDHCP-1. غيّر Display Name وhostname ليطابقا الجدول. افتح design/port_map.csv، ونفذ كل صف باستخدام طرفَي الوصلة والمنفذين المحددين. استخدام نوع الكابل Automatic مسموح في المختبر. انتظر وصول المنافذ إلى up؛ اللون وحده لا يثبت نجاح الخدمة.')
table(['الملف','استخدامه أثناء التطبيق'],[['design/port_map.csv','كل وصلة فعلية ومنفذَيها؛ المرجع عند توصيل 88 رابطًا'],['design/devices.csv','الموديل وRouter ID وعنوان إدارة كل جهاز'],['design/endpoints.csv','منفذ كل مستخدم وVLAN والعنوان الثابت أو DHCP'],['configs/<hostname>.txt','الإعداد الكامل لذلك الجهاز فقط'],['report/configuration_reference.docx','جميع ملفات الإعداد في مستند قابل للبحث والطباعة']])
subheading('الخطوة ٢ — إعداد اسم الجهاز والدخول الآمن')
paragraph('افتح الجهاز ثم CLI. عند سؤال initial configuration dialog أجب no، وانتظر اكتمال الإقلاع وظهور prompt. استخدم بداية ملف CORE-1 التالية مثالًا للإعداد المشترك. hostname يمنع الالتباس، وno ip domain-lookup يلغي انتظار DNS عند الخطأ في أمر، وdomain-name مطلوب لهوية مفتاح SSH. حساب netadmin محلي ومخصص للمختبر، وenable secret يحمي وضع الصلاحيات.')
code(config('CORE-1').split('interface Loopback0')[0].strip())
paragraph('VTY-IT قائمة تحكم في مصادر جلسات الإدارة تسمح فقط لـ10.50.3.10. transport input ssh يمنع Telnet؛ لا تغيّره إلى telnet ssh لتسهيل الأتمتة. إذا رفض إصدار المحاكي الصياغة أحادية السطر لـRSA، استخدم crypto key generate rsa ثم اختر 2048 عند السؤال، وبعدها تحقق من show ip ssh. لا تعتبر وجود الأمر بالنص وحده دليل نجاح جلسة SSH.')
subheading('الخطوة ٣ — إنشاء VLANs وتركيب منافذ Trunk')
paragraph('على كل Access أنشئ VLANs الخاصة بمبناه فقط وفق VLAN Matrix. مثال المبنى A التالي مأخوذ من ملف ACC-A1. VLAN 998 هي native غير مستخدمة للمستخدمين، و999 للمنافذ المغلقة. لا تضع عنوان IP على منفذ access للمستخدم؛ عضويته في VLAN هي التي تحدد شبكته.')
code('\n'.join(block('ACC-A1',f'vlan {v}') for v in [10,20,30,31,32,40,60,70,80,99,998,999]))
paragraph('على ACC-A1: Gi0/1 إلى DIST-A1 وGi0/2 إلى ACC-A2. يجب أن تتطابق native VLAN والقائمة المسموحة في طرفَي كل trunk. trust يخص DHCP Snooping ولا يعني فتح إدارة السويتش للمستخدمين. لا تطبق Port Security أو PortFast على هذه الوصلات.')
code(block('ACC-A1','interface GigabitEthernet0/1')+'\n'+block('ACC-A1','interface GigabitEthernet0/2'))
code('show vlan brief\nshow interfaces trunk\nshow spanning-tree')
paragraph('النتيجة المتوقعة: VLANs المطلوبة موجودة في allowed وكذلك allowed and active، وليس في allowed فقط. إذا ظهرت VLAN 99 وحدها active فلا تنتقل لاختبار DHCP؛ أصلح قاعدة VLAN أولًا. في التوليد الآلي احتجنا تمثيل ENGINE/VLANS الأصلي بجانب نص الإعداد.')
subheading('الخطوة ٤ — إعداد الوصلات الموجهة في Core')
paragraph('ip routing يمكّن التوجيه على 3650. no switchport يحول منفذ Core من Layer 2 إلى Layer 3. كل وصلة بين Core وDistribution لها /30 مستقل؛ المثال يصل CORE-1 إلى DIST-A1. يظل Loopback عنوانًا ثابتًا لاختيار Router ID. نفذ بقية الوصلات من ملف Core المخصص دون إعادة استخدام عنوان المثال.')
code(block('CORE-1','interface Loopback0')+'\n'+block('CORE-1','interface GigabitEthernet1/0/1'))
code(block('CORE-1','router ospf 1'))
paragraph('كل network مع wildcard 0.0.0.0 يحدد عنوان واجهة بعينه للإعلان في Area 0. passive-interface default يمنع الجيران افتراضيًا، ثم يُسمح بهم فقط على وصلات النقل. ip ospf network point-to-point يلائم وصلات الجارين ولا يتطلب انتخاب DR/BDR.')
subheading('الخطوة ٥ — إعداد التوزيع وHSRP وInter-VLAN Routing')
paragraph('على DIST-A1 أنشئ الواجهة الفيزيائية ثم native subinterface. كل VLAN لها واجهة Gi0/2.VLAN وقناعها الخاص. لا تستخدم عنوان VIP كعنوان حقيقي لواجهة. مثال VLAN 30: العنوان الحقيقي .2 على A1 و.3 على A2؛ كلاهما يشتركان في VIP .1. وجود ACL قبل ربطها مطلوب؛ انسخ تعريفاتها من ملف الجهاز قبل إتمام الربط.')
code(block('DIST-A1','interface GigabitEthernet0/2')+'\n'+block('DIST-A1','interface GigabitEthernet0/2.998'))
code(block('DIST-A1','interface GigabitEthernet0/2.30'))
code(block('DIST-A2','interface GigabitEthernet0/2.30'))
table(['الأمر/الحقل','سبب استخدامه'],[['encapsulation dot1Q 30','ربط الواجهة الفرعية بالإطارات الموسومة VLAN 30'],['standby 30 ip','عنوان البوابة الافتراضية الذي توزعه DHCP'],['priority 110 / 100','تفضيل A1 على A2 عندما يكون كلاهما متاحًا'],['preempt','إرجاع الدور للجهاز الأعلى أولوية بعد استعادته'],['ip helper-address','ترحيل طلبات DHCP إلى 10.50.2.138'],['ip access-group ... in','فحص حركة المستخدم الواردة قبل توجيهها']])
paragraph('كرر المبدأ لكل VLAN باستخدام عناوين الملف الخاص بالمبنى؛ لا تستخدم /24 لجميع الشبكات. الواجهة VLAN 99 للإدارة، وشبكة Server VLAN 50 موجودة في E. لم يُعتمد standby use-bia ولا HSRP tracking؛ يبقى التصميم وعنوان MAC الافتراضي القياسي كما هما.')
subheading('الخطوة ٦ — تشغيل OSPF على راوتر التوزيع')
code(block('DIST-A1','interface GigabitEthernet0/0')+'\n'+block('DIST-A1','interface GigabitEthernet0/1')+'\n'+block('DIST-A1','router ospf 1'))
paragraph('واجهات المستخدمين passive أصلًا؛ لا تضف جيران OSPF على VLANs المستخدمين. تحقق من جارَين FULL على Distribution واثني عشر على Core. ثم افحص مسار شبكة مبنى آخر ومسار 0.0.0.0. الجار FULL وحده لا يثبت أن سياسة ACL أو DNS تعمل.')
code('show ip interface brief\nshow ip ospf neighbor\nshow ip route\nshow ip route 0.0.0.0')
subheading('الخطوة ٧ — DHCP المركزي: عنوان الخادم والاستثناءات والنطاقات')
code(block('DHCP-1','interface GigabitEthernet0/0')+'\n'+'\n'.join(line for line in config('DHCP-1').splitlines() if line.startswith('ip route ') or line.startswith('ip dhcp excluded-address 10.10.0.'))+'\n'+block('DHCP-1','ip dhcp pool A_V30'))
paragraph('excluded-address يحجز أول 15 عنوانًا فلا تُمنح VIP أو عناوين الراوترات لمستخدم. network يحدد الشبكة والقناع، وdefault-router يحدد VIP، وdns-server يشير للخادم المركزي. يحتوي DHCP-1.txt على جميع النطاقات، وليس نطاق A فقط. الخادم متصل بمنفذ موثوق في ACC-E2؛ جميع منافذ العملاء غير موثوقة.')
paragraph('من PC افتح Desktop ثم IP Configuration واختر DHCP. أو من Command Prompt نفذ ipconfig /renew وانتظر النتيجة ثم ظهور محث أوامر PC قبل الأمر التالي. تحقق من IP والقناع والبوابة وDNS ثم طابق MAC مع show ip dhcp binding على DHCP-1. لا تُرسل ping بينما التجديد ما زال معلقًا، ولا تعتبر أول timeout حكمًا نهائيًا قبل استقرار ARP والتوجيه.')
code('ipconfig /renew\nipconfig\n')
code('show ip dhcp binding\nshow ip dhcp pool')
subheading('الخطوة ٨ — ضبط خوادم DNS والويب والمستخدم الإداري')
paragraph('على DNS-SERVER: Desktop ثم IP Configuration ثم Static؛ أدخل 10.50.2.139 والقناع 255.255.255.192 والبوابة 10.50.2.129. من Services ثم DNS اختر On وأضف سجلي A المحددين. على EDU-WEB استخدم 10.50.2.140 ونفس القناع والبوابة، ثم Services ثم HTTP ثم On. فعّل HTTPS إن دعمه الموديل؛ لم يُثبت اختبار HTTPS بعد. لا تشغّل DHCP إضافيًا على Server-PT.')
paragraph('على PUBLIC-WEB: العنوان 203.0.113.10 والقناع 255.255.255.0 والبوابة 203.0.113.1 مع HTTP On. على IT-ADMIN: 10.50.3.10/26 والبوابة 10.50.3.1. بقية أجهزة DHCP تتبع endpoints.csv، والكاميرات النموذجية والخوادم وIT بعناوين ثابتة خارج نطاقات الحجز الديناميكي.')
code('edu.campus.example       A   10.50.2.140\nwww.external.example    A   203.0.113.10')
paragraph('اختبر من متصفح الطالب http://edu.campus.example؛ نجاح تحميل الصفحة باسم النطاق دليل على معالجة الاسم والوصول عبر HTTP في تلك المحاولة. لمبة الوصلة أو ping الخادم ليست بديلًا لهذا الاختبار، خصوصًا أن ICMP إلى الخادم التعليمي محظور على الطلاب.')
subheading('الخطوة ٩ — ISP وEDGE وPAT')
code(config('ISP').strip())
paragraph('طبّق ملف EDGE.txt كاملًا. فيما يلي جزء النقل وNAT منه: المنفذان إلى Core هما inside والمنفذ إلى ISP هو outside. overload يسمح لعناوين المستخدمين بمشاركة عنوان EDGE الخارجي. المسار الافتراضي إلى ISP يُعلن داخل OSPF؛ لا تضع default على كل راوتر نحو عنوان غير متصل مباشرة.')
code(config('EDGE')[config('EDGE').index('interface Loopback0'):].strip())
code('show ip route 0.0.0.0\nshow ip nat translations\nshow ip nat statistics')
subheading('الخطوة ١٠ — الحفظ والمراجعة قبل اختبار الأعطال')
paragraph('احفظ الإعداد في كل جهاز بعد نجاح تطبيقه، ثم احفظ ملف Packet Tracer باسم جديد قبل أي تجربة تعطيل. لا تستخدم Fast Forward في جلسة هذا المشروع عند الحاجة لدليل استقرار، فقد التُقط انهيار داخل CMacTable::removeMacEntry أثناء معالجة مؤقتات المحاكاة. استخدم Realtime وانتظر prompt. لا يوجد دليل أن ترقية الإصدار وحدها حلت هذا العطل.')
code('end\nwrite memory\nshow running-config')
paragraph('يجب إعادة فتح الملف المحفوظ واختبار DHCP وOSPF وHSRP وHTTP وSSH وقواعد المنع، وليس الاكتفاء برسالة Save. ملف pkt المرفق حفظ أصلي بعد استعادة Gi0/2؛ أحدث صور الاتصال اللاحقة من جلسة تشخيصية منفصلة بنفس السياسة، ولم تُعتمد تلك الجلسة بديلًا عن الملف المحفوظ.')

heading('الفصل الخامس: أمن الشبكة')
table(['الفئة','المسموح','الممنوع / القيود'],[['الطلاب والمختبرات','DHCP وDNS وHTTP/HTTPS التعليمي والإنترنت المحاكى','HR والمالية والإدارة وكل الوجهات الداخلية الأخرى'],['الضيوف','DHCP وDNS المحدد والوجهات الخارجية','كل الشبكات الداخلية عدا استثناءات الخدمة المذكورة؛ لا وصول للويب التعليمي'],['هيئة التدريس','الخوادم وDNS وDHCP والإنترنت','HR والمالية والإدارة وبقية الشبكات الداخلية غير المصرح بها'],['الإدارة وشؤون الطلاب','الخوادم وHR والمالية والخدمات الداخلية والإنترنت','إدارة أجهزة الشبكة؛ ليست الإدارة الإدارية هي فريق IT'],['HR والمالية','الخدمات الداخلية والإنترنت وفق النموذج','إدارة الشبكة'],['Voice وSecurity','DNS والخدمة التعليمية المحددة للاختبار؛ DHCP للـVoice','بقية الشبكات الداخلية والإنترنت؛ لا ادعاء بتشغيل مكالمات أو NVR'],['IT-ADMIN','SSH إلى الأجهزة؛ وصول تشخيصي من VLAN 99','مصادر أخرى لا يجيزها VTY ACL']])
paragraph('ACLs موسعة مسماة، واردة على كل واجهة Layer 3 الفرعية غير إدارية ومتماثلة على جهازي HSRP. ترتيبها: رسائل HSRP من الجارين؛ DHCP broadcast وتجديد للخادم المحدد؛ استثناءات ردود البنية التحتية اللازمة؛ منع الإدارة والنقل وعناوين الأجهزة؛ السماح بالخدمات؛ منع النطاقات الداخلية؛ السماح الخارجي للدور المصرح؛ deny نهائي. لا يوجد permit ip any any غير مقيد بالمصدر.')
paragraph('تمنع السياسة الوصول المباشر لعناوين واجهات Layer 3 الفرعية بما فيها Ping إلى البوابة من شبكات المستخدمين، لكنها لا تمنع ARP أو مرور الحزم عبر البوابة إلى وجهة مسموحة. لذلك لا يصلح فشل Ping إلى Gateway كدليل على فشل التوجيه في هذا التصميم. اختبار الطالب للخدمة التعليمية يكون HTTP/HTTPS، وليس Ping الممنوع عمدًا.')
paragraph('ACL ليست جدار حماية ذا حالة: established يطابق أعلام TCP لا جلسة محفوظة. استثناء رد TCP من شبكة الخوادم إلى IT يسمح بإدارة DHCP-1؛ لا يوجد هذا الاستثناء للطلاب أو الضيوف. ACL على واجهة Layer 3 الفرعية لا تمنع اتصال جهازين داخل VLAN نفسها؛ لذلك عُزلت الأقسام الحساسة في VLANs مستقلة. لا ندعي حماية كاملة من انتحال IP داخل الشبكة.')
paragraph('إدارة الأجهزة عبر SSH v2 فقط، مستخدم netadmin بإعداد صلاحية 15؛ على 2960 لوحظ دخول User EXEC ثم الحاجة إلى enable، مع enable secret، login local، مهلة خمول، وتحذير دخول. VTY ACL تسمح فقط بعنوان IT-ADMIN. أجهزة Core لديها وصلات إدارة مستقلة إلى VLAN 99 في E. EDGE يستخدم Loopback99 بعنوان 10.255.254.30/32 يُوصل إليه من VLAN 99؛ هذه استثناء موثق لأن منافذ 2911 الثلاثة مستخدمة للنقل والخارج.')
paragraph('قيم LabOnly26Admin9 وLabOnly26Enable9 أمثلة مختبرية منشورة وليست أسرارًا حقيقية. تستبدل قبل أي استخدام فعلي ولا علاقة لها ببيانات الخادم السحابي. service password-encryption ليس تشفيرًا قويًا؛ الحماية الأساسية لكلمات المستخدم وenable هي secret. لا تُرفق مفاتيح SSH أو بيانات دخول السحابة بالتقرير.')
paragraph('Port Security على منافذ الأجهزة فقط: maximum 1 وsticky وviolation restrict. الـVoice ممثل بجهاز واحد دون PC خلف هاتف؛ عند تركيب هاتف وPC يجب تغيير الحد إلى 2 وتصميم voice VLAN ملائم. بعد التعلم يُحفظ running-config، وعند تغيير الجهاز تُزال MAC القديمة بشكل مقصود بدل تعطيل الحماية.')
paragraph('DHCP Snooping يعمل على Access للشبكات المحددة. uplinks إلى التوزيع موثوقة لأنها المسار لردود الخادم المركزي؛ منفذ DHCP-1 على ACC-E2 موثوق أيضًا. منافذ المستخدمين غير موثوقة ومحددة المعدل. Option 82 معطل في نموذج المحاكاة لتفادي اختلاف تعامل relay/server معه؛ يوثق القرار ولا يعمم على الإنتاج. المنافذ غير المستخدمة مغلقة ومربوطة بـVLAN 999.')
subheading('شرح عملي للأمان — تطبيق القواعد دون فتح صلاحيات زائدة')
paragraph('الخطوة ١: أنشئ ACL الموسعة في وضع الإعداد العام. مثال الطالب التالي هو التعريف الكامل A-V30-IN من ملف DIST-A1 دون اختصار. الترتيب مهم: DHCP وHSRP أولًا، ثم منع الإدارة وعناوين أجهزة الشبكة، ثم DNS والويب التعليمي، ثم منع الوجهات الخاصة غير المسموحة والسماح بالخارج للدور المصرح. deny الأخير يمنع أي حركة لم تطابق قاعدة سابقة.')
code(block('DIST-A1','ip access-list extended A-V30-IN'))
paragraph('الخطوة ٢: اربط القائمة بالواجهة Gi0/2.30 في الاتجاه in كما في الفصل الرابع، وكرر السياسة على A2 حتى لا يصبح failover مسارًا يتجاوز الحماية. لا تقلب ترتيب deny وpermit ولا تضف permit ip any any. افحص العداد قبل الاختبار وبعده إلى وجهة فعلية لها عنوان مؤكد من DHCP.')
code('show access-lists A-V30-IN\nshow ip interface GigabitEthernet0/2.30')
paragraph('الخطوة ٣: اضبط منفذ الطالب على ACC-A1. maximum 1 يسمح بعنوان مصدر واحد، وsticky يتعلم MAC، وrestrict يسقط المصدر المخالف ويزيد العداد؛ لا يشترط err-disabled. هذه السياسة ليست shutdown. PortFast وBPDU Guard للمستخدم النهائي فقط، وليس للوصلات بين سويتشات أو إلى راوتر HSRP.')
code(block('ACC-A1','interface FastEthernet0/3'))
paragraph('الخطوة ٤: فعّل DHCP Snooping للشبكات المحددة، واجعل trunk المار منه رد الخادم موثوقًا ومنفذ DHCP-1 موثوقًا على ACC-E2 فقط. حافظ على منافذ المستخدمين غير موثوقة. اختبار خادم مزيف يحتاج Server-PT بخدمة DHCP في نسخة اختبار مستقلة، لا PC عاديًا بلا خدمة DHCP.')
code('\n'.join(line for line in config('ACC-A1').splitlines() if line.startswith('ip dhcp snooping') or line.startswith('no ip dhcp snooping')))
paragraph('الخطوة ٥: المنافذ غير المستخدمة توضع في VLAN 999 وتغلق إداريًا. المثال التالي منفذ غير مستخدم على ACC-A1؛ راجع port_map قبل إغلاق أي منفذ حتى لا تفصل جهازًا مطلوبًا.')
code(block('ACC-A1','interface FastEthernet0/10'))
code('show port-security interface FastEthernet0/3\nshow ip dhcp snooping\nshow ip dhcp snooping binding\nshow interfaces status')
paragraph('الخطوة ٦: من IT-ADMIN فقط استخدم أمر SSH التالي بعنوان إدارة DIST-A1، ثم جرّب Telnet الذي يجب رفضه. المستخدم في هذه الملفات netadmin وليس admin. كلمة مرور المختبر موضحة في بداية الإعداد. وجود RSA في الحفظ يؤكد توليد المفتاح لكنه لا يغني عن تجربة وصول فعلية. لا تضف AAA تلقائيًا لمجرد ظهور > على 2960؛ يمكن استخدام enable مع توثيق السلوك المرصود.')
code('ssh -l netadmin 10.10.6.2\ntelnet 10.10.6.2')

heading('الفصل السادس: خطة الاختبار والأدلة')
paragraph(f'فحوص ساكنة منفذة: {validation["passed"]} نجحت، {validation["failed"]} فشلت. تشمل عدم التداخل وحدود الشبكات والسعة وهامش النمو وحجز العناوين وتفرد المنافذ وRouter IDs وتماثل HSRP ومطابقة إعدادات ACL وتقييم first-match لحزم اختبار ممثلة. هذه لا تقيس تشغيل IOS أو زمن convergence أو نجاح DHCP داخل المحاكي.')
showcmds=['show vlan brief','show interfaces trunk','show ip interface brief','show ip route','show ip ospf neighbor','show ip ospf','show standby','show access-lists','show ip dhcp binding','show ip dhcp snooping','show ip dhcp snooping binding','show port-security','show port-security interface FastEthernet0/1','show running-config']
scenarios=[('T01','VLAN/Trunk','تحقق من VLAN IDs وallowed list وnative 998','الشبكات المطلوبة فقط على كل وصلة'),('T02','DHCP طالب وموظف','DHCP على PC طالب وPC موظف؛ ipconfig /all؛ show binding','عنوان ضمن نطاق الشبكة وVIP وDNS صحيحان'),('T03','OSPF','show ip ospf neighbor وshow ip route على Core وDIST','جيران Full ومسارات داخلية؛ default من EDGE'),('T04','طالب → Web','فتح http://edu.campus.example من طالب','نجاح HTTP وDNS'),('T05','طالب → HR/Finance','TCP أو ICMP إلى مضيف فعلي في VLAN 12 و13','منع مع زيادة عداد ACL المناسب'),('T06','ضيف → الداخل','محاولة HTTP للسيرفر وSSH للإدارة','منع؛ مع استمرار DNS والخارج'),('T07','إدارة → HR/Finance','اختبار من VLAN 10 إلى مضيفَي HR/Finance','اتصال مسموح'),('T08','SSH فقط','من IT-ADMIN نفذ SSH للأجهزة ثم حاول Telnet','SSH ناجح؛ Telnet مرفوض؛ مصدر غير IT مرفوض'),('T09','HSRP تعطل جهاز','سجل show standby ثم أغلق Gi0/2 في DIST-A1 أثناء اختبار اتصال خارجي','A2 Active وعودة الاتصال؛ تسجيل عدد الحزم المفقودة'),('T10','HSRP استعادة','أعد Gi0/2 باستخدام no shutdown وانتظر HSRP','عودة A1 Active بسبب preempt'),('T11','تعطل رابط واحد','أغلق وصلة Core واحدة من A1 ثم أعدها؛ لا يوجد HSRP tracking في النسخة المعتمدة','OSPF يستخدم الوصلة الأخرى؛ لا يشترط تغير دور HSRP'),('T12','DHCP مزيف','شبكة اختبار فقط: صل Server DHCP على منفذ مستخدم غير موثوق','عدم قبول OFFER المزيف واستمرار الحصول من الخادم المعتمد'),('T13','Port Security','تعلم MAC ثم استبدل جهاز المنفذ بجهاز مختلف','زيادة violation/drop وفق restrict؛ استعادة الجهاز الأصلي'),('T14','منافذ مغلقة','وصل جهازًا بمنفذ غير مستخدم','المنفذ administratively down وVLAN 999'),('T15','الإنترنت وPAT','افتح www.external.example وافحص show ip nat translations','HTTP ناجح وترجمة مصدر ظاهرة'),('T16','الحفظ والاستعادة','write memory ثم حفظ pkt وإغلاقه وفتحه','بقاء الطوبولوجيا والإعدادات وإعادة تحقق الخدمات')]
table(['ID','الاختبار','الطريقة','النتيجة المتوقعة'],scenarios)
table(['ID','الحالة الفعلية','الملاحظة','الدليل'],[[row[0],runtime.get(row[0],{}).get('status','PENDING'),runtime.get(row[0],{}).get('actual','لم يُنفذ بعد'),runtime.get(row[0],{}).get('evidence_file','')] for row in scenarios])
paragraph('لا توجد نتائج تشغيلية مختلقة في هذا الفصل. كل حالة غير مدعومة بنتيجة ودليل تبقى PENDING. في T09 يمكن استخدام ping إلى 203.0.113.10 من جهاز طالب؛ لا نستخدم Ping إلى البوابة أو خادم التعليم لأن سياسة ACL تمنعهما. Packet Tracer قد لا يدعم ping -t على PC؛ تُستخدم طلبات متكررة أو Complex PDU دوري، ويسجل الأسلوب المستخدم.')
paragraph('توثيق كل اختبار: اسم الجهاز، الإعداد الأولي، الإجراء، النتيجة المتوقعة، الفعلية، وقت الاختبار، واسم لقطة الشاشة. لا تعد صورة المخطط دليلًا على نجاح التوجيه أو فشل/نجاح ACL. تعرض لقطة failover الجهاز Active قبل وبعد مع دليل مرور الخدمة.')
table(['أوامر التحقق المطلوبة'],[[c] for c in showcmds])
heading('الفصل السابع: استكشاف الأعطال')
table(['العرض','طريقة الكشف','سبب محتمل','المعالجة'],[['لا يحصل PC على IP','ipconfig + show ip dhcp binding/snooping','VLAN خاطئة أو helper مفقود أو scope غير صحيح','طابق المنفذ والقناع والـrelay والمنفذ الموثوق ثم جدّد الطلب'],['HSRP جهازان Active','show standby + show interfaces trunk','انقطاع VLAN/Trunk بين الجارين أو ACL تحجب hello','استعد مسار Layer 2 وقارن group/VIP وسمح HSRP'],['لا يظهر جار OSPF','show ip ospf neighbor + interface brief','وصلة down أو passive أو area/network مختلفة','صحح الواجهة والمسار ونطاق إعلان OSPF'],['Ping البوابة يفشل والخدمة تعمل','ACL counters + HTTP للخدمة','منع مقصود لعناوين إدارة الأجهزة','لا تغيّر السياسة لإصلاح سلوك مقصود؛ اختبر الوجهة المسموحة'],['الضيف يصل لخادم داخلي','show access-lists + واجهة Layer 3 الفرعية config','permit واسع قبل deny أو ACL غير مطبقة على standby','صحح الترتيب والاتجاه على الجارين وأعد اختبار الحظر'],['بعد تبديل PC لا يمر المرور','show port-security interface','Sticky MAC لجهاز سابق','أعد الجهاز أو امسح MAC القديمة بإجراء تغيير معتمد'],['فشل DNS مع نجاح IP','DNS service + ipconfig + ACL counters','DNS غير مفعّل أو سجل مفقود أو ACL تحجب TCP/UDP53','فعّل الخدمة والسجلات والقواعد المحددة'],['تعطل DHCP رغم نجاح HSRP','فحص DHCP-1 ومساره ووصلته','الخادم المركزي أو ACC-E2 نقطة فشل','استعد الخدمة؛ لا تدّع أن HSRP يحمي خادم DHCP'],['لا تعمل ميزة في PT','سجل رسالة Invalid input وإصدار/موديل الجهاز','قيد محاكاة أو صياغة غير مدعومة','جرّب الموديل الصحيح؛ وثق القيد ولا تختلق مخرجات'],['الإنترنت المحاكى لا يعمل','show ip route وshow ip nat translations','غياب default أو NAT inside/outside أو ضبط ISP','صحح EDGE وISP والخادم ثم أعد اختبار HTTP']])
paragraph('أعطال فعلية: فقد قاعدة VLAN الأصلية منع DHCP رغم وجود VLAN في نص الإعداد؛ أُصلح بإضافة قاعدة VLAN في ملف المحاكي. فقد ارتباط ACL على SVI عولج باعتماد واجهات 2911 الفرعية. التقط GDB انهيار SIGSEGV داخل CMacTable::removeMacEntry عبر CScheduler وfastForwardTime؛ السبب الجذري غير محسوم، ولا يصح نسبته إلى المتصفح وحده. لاحقًا نجح Ping الطالب للخارج 4/4 وتحميل صفحة التعليم باسم النطاق في جلسة مستقرة دون تسريع. تظل اختبارات الحماية الشاملة والاستقرار النهائي غير مكتملة.')
paragraph('هذه أمثلة تشخيصية متوقعة وليست سجل أعطال تم افتعالها أو قياسها. سجل التشغيل الفعلي مستقل. إجراءات الرجوع: حفظ نسخة pkt قبل التغيير، نسخ running-config، إعادة no shutdown بعد تجارب التعطل، وإعادة DHCP/المنافذ لحالتها الأصلية بعد الاختبارات الأمنية.')
heading('الملاحق: التسليم وقابلية إعادة الإنتاج')
paragraph('الملفات: design/campus.json المصدر الموحد؛ addressing_plan.xlsx وCSV؛ configs لكل جهاز؛ diagrams/topology.png وtopology.drawio؛ pt/build_steps.md؛ tests/test_matrix.csv وvalidation_results.json؛ report/report.docx وPDF عند توليده. pt/university.pkt ولقطات المحاكي لا تعتبر موجودة أو ناجحة إلا إذا ذُكرت حالتها في manifest التسليم.')
paragraph('إعادة التوليد: python tools/build_project.py ثم python tools/validate_project.py ثم python tools/build_documents.py. أي تغيير في العناوين يجب أن يبدأ من المصدر/المولد، لا من تعديل جدول Word فقط. مولد الوثائق يحتفظ بالنتائج من tests/runtime_results.json؛ لا يستبدلها بنتائج متوقعة أو بقيم ناجحة تلقائية.')
paragraph('المراجع: project.pdf الصفحات 1–14؛ تحليل سيناريو الجامعة وتصميم هندسي موثق؛ ملفات الإعداد والتحقق المرفقة. لا تستخدم لقطات خارجية أو نتائج من مشاريع أخرى بوصفها دليلًا لهذا المشروع.')

# Marked cover placeholders, never invented identities.
cover={'student':'[اسم الطالب]','student_id':'[الرقم الجامعي]','university':'[اسم الجامعة]','instructor':'[اسم المدرس]'}
if (ROOT/'report/cover.json').exists(): cover.update(json.loads((ROOT/'report/cover.json').read_text()))
else: out('report/cover.json',json.dumps(cover,indent=2,ensure_ascii=False)+'\n')
doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.27); sec.page_height=Inches(11.69); sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
style=doc.styles['Normal']; style.font.name='DejaVu Sans'; style.font.size=Pt(10)
style.paragraph_format.space_after=Pt(7)
style.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.RIGHT
style.element.get_or_add_pPr().append(OxmlElement('w:bidi'))
settings=doc.settings.element;lang=OxmlElement('w:themeFontLang');lang.set(qn('w:val'),'ar-SA');lang.set(qn('w:bidi'),'ar-SA');settings.append(lang)
for sty in ['Title','Heading 1','Heading 2']:
 doc.styles[sty].font.name='DejaVu Sans'; doc.styles[sty].font.color.rgb=RGBColor.from_string('17324D')
def rtl(p):
 p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 pr=p._p.get_or_add_pPr()
 if pr.find(qn('w:bidi')) is None: pr.append(OxmlElement('w:bidi'))
 for run in p.runs:
  rp=run._r.get_or_add_rPr();lang=OxmlElement('w:lang');lang.set(qn('w:val'),'ar-SA');lang.set(qn('w:bidi'),'ar-SA');rp.append(lang)
  fonts=rp.find(qn('w:rFonts'))
  if fonts is None: fonts=OxmlElement('w:rFonts');rp.append(fonts)
  fonts.set(qn('w:cs'),'DejaVu Sans')
 return p

def add_code(target,text,size=8):
 # Right-aligned on the Arabic page, but the IOS character order stays LTR.
 lines=text.splitlines()
 for start in range(0,len(lines),28):
  p=target.add_paragraph('\n'.join(lines[start:start+28]));p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
  pf=p.paragraph_format;pf.space_before=Pt(2);pf.space_after=Pt(4);pf.line_spacing=1;pf.keep_together=False
  pr=p._p.get_or_add_pPr();direction=OxmlElement('w:bidi');direction.set(qn('w:val'),'0');pr.append(direction)
  for run in p.runs:
   run.font.name='DejaVu Sans Mono';run.font.size=Pt(size)
   el=OxmlElement('w:rtl');el.set(qn('w:val'),'0');run._r.get_or_add_rPr().append(el)
p=rtl(doc.add_paragraph('تصميم وتنفيذ شبكة جامعة\nمتعددة الكليات والمباني','Title'))
rtl(doc.add_paragraph('Network Analysis and Design\nالتكليف النهائي العملي'))
for label,key in [('الجامعة','university'),('الطالب','student'),('الرقم الجامعي','student_id'),('إشراف','instructor')]: rtl(doc.add_paragraph(label+': '+cover[key]))
rtl(doc.add_paragraph('نسخة التصميم والإعداد — حالة الاختبارات التشغيلية موضحة في الفصل السادس'))
rtl(doc.add_paragraph('التاريخ: 2026-09-13'))
doc.add_page_break(); doc.add_picture(str(ROOT/'diagrams/topology.png'),width=Inches(7))
rtl(doc.add_paragraph('الشكل 1: مخطط هندسي مولّد؛ ليس لقطة من المحاكي.'))
markdown=['# تصميم وتنفيذ شبكة جامعة متعددة الكليات والمباني\n']
for kind,data in sections:
 if kind=='h':
  hp=rtl(doc.add_heading(data,level=1))
  if data.startswith('الفصل'): hp.paragraph_format.page_break_before=True
  markdown.append('## '+data+'\n')
 elif kind=='h2':
  rtl(doc.add_heading(data,level=2));markdown.append('### '+data+'\n')
 elif kind=='code':
  add_code(doc,data);markdown.append('```text\n'+data+'\n```\n')
 elif kind=='p': rtl(doc.add_paragraph(data)); markdown.append(data+'\n')
 else:
  headers,rows=data; t=doc.add_table(rows=1,cols=len(headers)); t.style='Light Shading Accent 1'
  t.alignment=WD_TABLE_ALIGNMENT.RIGHT
  prop=t._tbl.tblPr; bidi=OxmlElement('w:bidiVisual'); prop.append(bidi)
  for i,h in enumerate(headers): t.rows[0].cells[i].text=str(h)
  header_props=t.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); header_props.append(repeat)
  for row in rows:
   cells=t.add_row().cells
   for i,value in enumerate(row): cells[i].text=str(value)
  for row in t.rows:
   for cell in row.cells:
    for p in cell.paragraphs:
     rtl(p)
     for run in p.runs: run.font.size=Pt(7.5)
  markdown.append(mdtable(headers,rows))
rtl(doc.add_heading('ملحق صور التشغيل الفعلية',level=1))
rtl(doc.add_paragraph('هذه قصاصات مباشرة من نوافذ المحاكي، دون تحسين OCR أو تركيب مخرجات. الصور 01 و02 في المجلد من النموذج الأولي السابق ولا تعتمد لإثبات النسخة الحالية. نجاح HSRP control-plane لا يثبت استمرارية الخدمة.'))
doc.add_picture(str(ROOT/'screenshots/04_access_trunk.png'),width=Inches(6.4))
rtl(doc.add_paragraph('VLANs المسموحة والنشطة على ACC-A1'))
markdown.append('![VLANs المسموحة والنشطة على ACC-A1](../screenshots/04_access_trunk.png)\n')
doc.add_picture(str(ROOT/'screenshots/05_student_renew.png'),width=Inches(6.4))
rtl(doc.add_paragraph('تجديد DHCP للطالب: 10.10.0.17؛ نتيجة التجديد أسفل الشاشة'))
markdown.append('![تجديد DHCP للطالب: 10.10.0.17؛ نتيجة التجديد أسفل الشاشة](../screenshots/05_student_renew.png)\n')
doc.add_picture(str(ROOT/'screenshots/06_admin_dhcp.png'),width=Inches(6.4))
rtl(doc.add_paragraph('DHCP للإدارة: 10.10.5.145/26 والبوابة 10.10.5.129'))
markdown.append('![DHCP للإدارة: 10.10.5.145/26 والبوابة 10.10.5.129](../screenshots/06_admin_dhcp.png)\n')
doc.add_picture(str(ROOT/'screenshots/07_dhcp_bindings.png'),width=Inches(6.4))
rtl(doc.add_paragraph('عقود DHCP الفعلية على DHCP-1؛ الصفحة الأولى'))
markdown.append('![عقود DHCP الفعلية على DHCP-1؛ الصفحة الأولى](../screenshots/07_dhcp_bindings.png)\n')
doc.add_picture(str(ROOT/'screenshots/08_hsrp_baseline.png'),width=Inches(6.4))
rtl(doc.add_paragraph('DIST-A1 Active وDIST-A2 Standby قبل العطل'))
markdown.append('![DIST-A1 Active وDIST-A2 Standby قبل العطل](../screenshots/08_hsrp_baseline.png)\n')
doc.add_picture(str(ROOT/'screenshots/09_hsrp_failover.png'),width=Inches(6.4))
rtl(doc.add_paragraph('DIST-A2 Active بعد shutdown لواجهة LAN في A1'))
markdown.append('![DIST-A2 Active بعد shutdown لواجهة LAN في A1](../screenshots/09_hsrp_failover.png)\n')
doc.add_picture(str(ROOT/'screenshots/10_hsrp_preemption.png'),width=Inches(6.4))
rtl(doc.add_paragraph('عودة A1 إلى Active بعد no shutdown بسبب preempt'))
markdown.append('![عودة A1 إلى Active بعد no shutdown بسبب preempt](../screenshots/10_hsrp_preemption.png)\n')
doc.add_picture(str(ROOT/'screenshots/11_external_ping_failed.png'),width=Inches(6.4))
rtl(doc.add_paragraph('محاولة Ping خارجية سابقة فاشلة؛ تلتها محاولة ناجحة موثقة في الملحق'))
markdown.append('![محاولة Ping خارجية سابقة فاشلة؛ تلتها محاولة ناجحة موثقة في الملحق](../screenshots/11_external_ping_failed.png)\n')
rtl(doc.add_paragraph('الأدلة التالية أحدث من محاولة الفشل السابقة. جُمعت في جلسة تشخيصية مستقرة بنفس سياسة العناوين وACL؛ محاولة أمر MAC aging غير المدعوم لم تُعتمد. لا تعني هذه الصور أن جميع الاختبارات اجتازت أو أن ملف pkt المرفق أُعيد اختباره نهائيًا.'))
doc.add_picture(str(ROOT/'screenshots/12_student_external_success.png'),width=Inches(6.4))
rtl(doc.add_paragraph('محاولة لاحقة: اتصال الطالب بالخادم الخارجي 4/4 بعد الاستقرار'))
markdown.append('![محاولة لاحقة: اتصال الطالب بالخادم الخارجي 4/4 بعد الاستقرار](../screenshots/12_student_external_success.png)')
doc.add_picture(str(ROOT/'screenshots/13_student_educational_http.png'),width=Inches(6.4))
rtl(doc.add_paragraph('صفحة التعليم محملة من متصفح الطالب باستخدام اسم النطاق'))
markdown.append('![صفحة التعليم محملة من متصفح الطالب باستخدام اسم النطاق](../screenshots/13_student_educational_http.png)')
doc.add_picture(str(ROOT/'screenshots/14_current_core_ospf.png'),width=Inches(6.4))
rtl(doc.add_paragraph('اثنا عشر جار OSPF بحالة FULL على Core في التصميم الحالي'))
markdown.append('![اثنا عشر جار OSPF بحالة FULL على Core في التصميم الحالي](../screenshots/14_current_core_ospf.png)')
doc.add_picture(str(ROOT/'screenshots/15_student_port_security.png'),width=Inches(6.4))
rtl(doc.add_paragraph('منفذ الطالب Secure-up؛ restrict؛ صفر مخالفات في الفحص الأساسي'))
markdown.append('![منفذ الطالب Secure-up؛ restrict؛ صفر مخالفات في الفحص الأساسي](../screenshots/15_student_port_security.png)')
doc.add_picture(str(ROOT/'screenshots/16_student_nat.png'),width=Inches(6.4))
rtl(doc.add_paragraph('ترجمات NAT فعلية لحركة الطالب؛ لا تثبت HTTP الخارجي بمفردها'))
markdown.append('![ترجمات NAT فعلية لحركة الطالب؛ لا تثبت HTTP الخارجي بمفردها](../screenshots/16_student_nat.png)')
footer=doc.sections[0].footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('University Campus Network | ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.save(ROOT/'report/report.docx'); out('report/report.md','\n'.join(markdown))

reference=Document();rs=reference.sections[0];rs.page_width=Inches(8.27);rs.page_height=Inches(11.69);rs.left_margin=Inches(.6);rs.right_margin=Inches(.6)
reference.styles['Normal'].font.name='DejaVu Sans';reference.styles['Normal'].font.size=Pt(10)
rtl(reference.add_heading('مرجع الإعدادات الكامل لجميع الأجهزة',level=0))
rtl(reference.add_paragraph('هذا الملحق يحتوي النص الكامل لملفات configs للأجهزة الخمسة والعشرين، دون اختصار أو حذف لقواعد ACL. الشرح خطوة بخطوة في الفصلين الرابع والخامس من التقرير الرئيسي. هذه إعدادات تصميمية وليست مخرجات show running-config مقتبسة من المحاكي.'))
rtl(reference.add_paragraph('استخدم ملف TXT الخاص بالجهاز للنسخ إلى CLI؛ لا تلصق جميع الأجهزة في جلسة واحدة. الكود محافظ على ترتيب حروفه اللاتيني ومحاذاته إلى يمين الصفحة، والشرح والعناوين عربية RTL. كلمات المرور أمثلة مختبرية منشورة فقط.'))
for n in m['nodes']:
 name=n['name'];hp=rtl(reference.add_heading(name,level=1));hp.paragraph_format.page_break_before=True
 rtl(reference.add_paragraph('الجهاز: '+name+' — ملف النسخ: configs/'+name+'.txt'))
 add_code(reference,config(name).strip(),size=7.5)
reference.save(ROOT/'report/configuration_reference.docx')

out('design/addressing_plan.md','# خطة التصميم والعنونة\n\nالأعداد افتراضية؛ نمو 25% و15 عنوانًا محجوزًا. أول مضيف VIP، التاليان لجهازي التوزيع.\n\n'+mdtable(['Building/VLAN','Name','Assumed','Growth','Network','Mask','First/Gateway','Last','Broadcast','Usable'],[[f'{s["building"]}/{s["vlan"]}',s['name'],s['assumed_endpoints'],s['growth_endpoints'],s['network'],s['mask'],s['first_host'],s['last_host'],s['broadcast'],s['usable']] for s in ss]))
out('tests/show_commands.txt','\n'.join(showcmds)+'\n\nAdditional diagnostics:\nshow spanning-tree\nshow ip nat translations\nshow ip dhcp pool\nshow version\n')
with (ROOT/'tests/test_matrix.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f); w.writerow(['id','test','procedure','expected','actual','status','evidence_file']); w.writerows([list(row)+[runtime.get(row[0],{}).get('actual',''),runtime.get(row[0],{}).get('status','PENDING'),runtime.get(row[0],{}).get('evidence_file','')] for row in scenarios])
out('screenshots/README.md','# أدلة التشغيل\n\nلقطات 01 و02 تخص النموذج الأولي السابق ذي 98 وصلة وليست دليل النسخة الحالية. صورة 03 هي DHCP من التشغيل السابق للتصميم الحالي. الصور 04–11 من جلسة استعادة 2026-09-13، وترتبط بالاختبارات في tests/runtime_results.json. الصورة 11 تسجل محاولة فاشلة سابقة. الصور 12–16 من جلسات لاحقة: Ping الطالب ناجح وHTTP تعليمي وOSPF وPort Security أساسي وNAT. ليست أدلة على اكتمال كل اختبارات الأمن. جميع الصور قصاصات بكسلات أصلية دون OCR. حالة المشروع: تحقق جزئي لا تسليم نهائي مكتمل. صورة diagrams/topology.png مخطط تصميم وليست دليل تشغيل. لا تُدرج بيانات دخول حساب Cisco أو مفاتيح الخادم في الصور.\n')
out('pt/build_steps.md',f'''# دليل بناء المحاكاة خطوة بخطوة

## الحالة
ملف university.pkt محفوظ من Packet Tracer 9 بعد تجربة HSRP وإعادة Gi0/2 إلى العمل. التحقق جزئي؛ نجح HTTP التعليمي وPing خارجي لاحقًا في جلسة تشخيصية، لكن اختبارات الحماية واستمرارية الخدمة أثناء العطل وإعادة الفتح النهائية غير مكتملة. بعد الفتح انتظر التقارب ثم جدد DHCP باستخدام ipconfig /renew عند الحاجة. university.xml مصدر مولد قبل التشغيل وليس نسخة مفكوكة من الحفظ الأخير. يحتوي الحفظ الأصلي على 71 جهاز شبكة/طرفية و50 كائن Power Distribution أضافها المحاكي؛ هذه ليست أجهزة شبكة إضافية.

## 1. الأجهزة
- 2 × 3650-24PS: CORE-1/2.
- 10 × 2911: DIST-A1/A2 إلى DIST-E1/E2 (واجهات فرعية 802.1Q).
- 10 × 2960-24TT: ACC-A1/A2 إلى ACC-E1/E2.
- 3 × 2911: EDGE وISP وDHCP-1.
- 46 جهازًا تمثيليًا حسب design/endpoints.csv؛ Server-PT للويب وDNS والخارج، PC-PT للبقية.
- عدد الأجهزة الطرفية يمثل عينة تعليمية؛ أعداد VLSM أكبر للتخطيط المستقبلي.

## 2. التوصيل
اتبع كل صف في design/port_map.csv حرفيًا (88 رابطًا). استخدم نوع الكابل الملائم أو Automatic، ثم تحقق من أرقام المنافذ. لا توصل VLANs بين المباني عبر Layer 2.

## 3. تطبيق الإعداد
ابدأ Core ثم التوزيع ثم الوصول ثم DHCP-1 وEDGE وISP. أجب no عن Initial Configuration Dialog. ألصق الملف المناسب من configs/ في CLI من وضع User EXEC. راجع كل Invalid input. احفظ write memory.
قد يحتاج RSA في بعض نسخ المحاكاة صياغة `crypto key generate rsa` ثم اختيار 2048 عند الطلب بدل الصياغة أحادية السطر. تحقق من `show ip ssh` ولا تفترض توليد المفتاح لمجرد لصق الأمر.

## 4. إعداد الخدمات
DHCP-1 معد بالكامل عبر CLI. لا تستخدم DHCP إضافيًا على Server-PT بالتوازي.
- DNS-SERVER = {services['dns']}: Static mask/gateway حسب design/endpoints.csv، Services > DNS > On.
- أضف edu.campus.example → {services['educational_web']}.
- أضف www.external.example → {services['public_web']}.
- EDU-WEB: Static IP ثم Services > HTTP > On. صفحة تعريف الجامعة بسيطة تكفي.
- PUBLIC-WEB: IP {services['public_web']} /24 وGateway 203.0.113.1 وHTTP On.
- IT-ADMIN = {services['it_admin']}: Static حسب الجدول.
- كل endpoint مكتوب له DHCP: Desktop > IP Configuration > DHCP. لا تعطه IP ثابتًا داخل النطاق الديناميكي.
- أجهزة Security وDNS/Web وIT ثابتة حسب الجدول.

## 5. اختبار صغير قبل التوسيع
نفذ مبنى A مع Core وخدمات E أولًا. تحقق من HSRP وOSPF وDHCP Snooping وPort Security والمنافذ الفعلية. إذا غاب دعم ميزة عن الجهاز، وثق رسالة الخطأ واختر موديلًا يدعمها؛ لا تستبدل HSRP بـVRRP وتسميه مكافئًا للتكليف.

## 6. التحقق الأمني
من طالب: HTTP للتعليم يجب أن يعمل؛ Ping البوابة ممنوع عمدًا. من ضيف: DNS والخارج فقط. من IT: SSH للأجهزة. من إدارة: HR/Finance مسموح. اتبع tests/test_matrix.csv وسجل النتائج الفعلية.

## 7. HSRP Failover
سجل show standby قبل التعطل. شغّل Ping متكرر إلى {services['public_web']} من A طالب أو Complex PDU دوري. أطفئ DIST-A1 من Physical أو أغلق الواجهة المناسبة، ثم سجل A2 Active وعدد الحزم المفقودة وعودة الاتصال. أعد A1 وتحقق من preempt. لا تعدّل ACL لتزييف نجاح Ping إلى وجهة محظورة.

## 8. حفظ التسليم
احفظ pt/university.pkt. أغلق وافتح الملف من جديد وكرر DHCP وOSPF وHSRP وHTTP/SSH والمنع. التقط أوامر tests/show_commands.txt من الأجهزة المناسبة؛ DHCP binding من DHCP-1 وSnooping/Port Security من Access.

## 9. قيود معلنة
الخادم DHCP-1 وACC-E2 وEDGE نقاط فشل منفردة. Voice يمثل عزلًا شبكيًا لا مكالمات. ACL لا تعزل مستخدمين داخل VLAN نفسها. تخصيصات الإنتاج تحتاج حصر أجهزة حقيقيًا ومراجعة منافذ/PoE/مسافات. كلمات الدخول داخل configs تجريبية منشورة وتغيّر للإنتاج.
''')
print('Generated Arabic DOCX/Markdown, Excel, PNG/drawio and build/test guides.')
